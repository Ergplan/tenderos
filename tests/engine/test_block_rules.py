"""Time-block supply floors in the sizing LP (SECI RTC-TM, Amendment-01 cl. 8.1): a floor in every
hour of the peak hours and of the off-peak hours, hard or with a shortfall penalty of
penaltyMultiple x the year's tariff per kWh short."""
from __future__ import annotations

import numpy as np
import pytest

from engine.sizing import rtc_lp as M
from tests.engine.test_bid_lp import BESS0, FIN0, HOUR
from tests.engine.test_rtc_lp import COSTS, _case

pytestmark = pytest.mark.skipif(M.highspy is None, reason="highspy not installed")

PEAK = ((HOUR >= 18) & (HOUR < 24)).astype(float).tolist()  # six evening hours


def _rules(mult=None):
    extra = {} if mult is None else {"penaltyMultiple": mult}
    return [
        {"id": "peak", "label": "Peak blocks", "target": 0.9, "hours": "peak", "basis": "block", **extra},
        {"id": "offpeak", "label": "Off-peak blocks", "target": 0.7, "hours": "offpeak", "basis": "block", **extra},
    ]


def _payload(rules, **extra):
    solar, wind, _ = _case()
    payload = {
        "ctx": {"demand": [100.0] * M.HOURS, "solarCf": solar.tolist(), "windCf": wind.tolist(), "plantMw": 100},
        "bess": BESS0, "costs": COSTS, "fin": FIN0, "dfrTarget": 0.0, "dfrBasis": "annual",
        "compliance": rules, "peakMask": PEAK, "tariffGuess": 6.0, "returnHourly": True, "returnLifetime": True,
    }
    payload.update(extra)
    return payload


@pytest.fixture(scope="module")
def hard():
    return M.solve(_payload(_rules()))


def test_hard_block_floors_hold_in_every_hour(hard):
    h = hard["hourly"]
    delivered = np.asarray(h["direct"]) + np.asarray(h["discharge"])
    floor = np.where(np.asarray(PEAK) > 0.5, 90.0, 70.0)
    assert (delivered >= floor - 0.02).all(), "every hour at or above its floor (hourly values are rounded to 0.01 MW)"
    for rule in hard["perYear"][0]["rules"]:
        assert rule["basis"] == "block" and rule["met"] and rule["hoursMet"] == 1.0
        assert rule["shortfallMu"] == pytest.approx(0, abs=1e-6)
    assert hard["perYear"][0]["blockShortfallMu"] == pytest.approx(0, abs=1e-6)
    assert all(row["blockShortfallMwh"] == pytest.approx(0, abs=1e-3) for row in hard["lifetime"])


@pytest.mark.slow  # three more solves: run with `make test-slow`
def test_a_penalty_lets_the_plant_fall_short_when_that_is_cheaper(hard):
    free = M.solve(_payload(_rules(0.0)))  # no penalty: shortfall costs nothing
    priced = M.solve(_payload(_rules(1.5)))  # SECI: 1.5 x tariff on the energy short
    steep = M.solve(_payload(_rules(50.0)))  # dearer than any plant: back to the hard floors
    tol = 0.003  # HiGHS gap tolerance
    assert free["tariff"] <= priced["tariff"] + tol <= hard["tariff"] + 2 * tol
    assert free["perYear"][0]["blockShortfallMu"] > priced["perYear"][0]["blockShortfallMu"] - 1e-6
    # a steep penalty leaves only the few hours dearer to cover than ~Rs 300/kWh short
    floor_mu = (0.9 * sum(PEAK) + 0.7 * (M.HOURS - sum(PEAK))) * 100 / 1000
    assert steep["perYear"][0]["blockShortfallMu"] <= min(priced["perYear"][0]["blockShortfallMu"] + 1e-6, 0.001 * floor_mu)
    assert priced["tariff"] - tol <= steep["tariff"] <= hard["tariff"] + tol
    # the reported shortfall is the hourly gap below each hour's floor
    for res in (free, priced):
        h = res["hourly"]
        delivered = np.asarray(h["direct"]) + np.asarray(h["discharge"])
        floor = np.where(np.asarray(PEAK) > 0.5, 90.0, 70.0)
        gap = np.maximum(0.0, floor - delivered).sum() / 1000
        assert res["perYear"][0]["blockShortfallMu"] == pytest.approx(gap, abs=0.05)
        rules = {r["id"]: r for r in res["perYear"][0]["rules"]}
        assert rules["peak"]["shortfallMu"] + rules["offpeak"]["shortfallMu"] == pytest.approx(res["perYear"][0]["blockShortfallMu"], abs=1e-6)
        assert rules["peak"]["penaltyMultiple"] in (0.0, 1.5)


@pytest.mark.parametrize("rules, extra, message", [
    ([{"id": "a", "target": 0.8, "hours": "all", "basis": "annual", "penaltyMultiple": 1.5}], {}, "penaltyMultiple"),
    ([{"id": "a", "target": 0.8, "hours": "all", "basis": "block", "penaltyMultiple": -1}], {}, "penaltyMultiple"),
    ([{"id": "a", "target": 0.8, "hours": "offpeak", "basis": "block"}], {"peakMask": [0.0] * M.HOURS}, "peakMask"),
    ([{"id": "a", "target": 0.8, "hours": "all", "basis": "daily"}], {}, "basis"),
])
def test_block_rule_inputs_are_checked(rules, extra, message):
    with pytest.raises(M.LpInputError, match=message):
        M.solve(_payload(rules, **extra))


def test_fifteen_minute_dispatch_holds_the_block_floors_and_writes_the_workbook(tmp_path):
    """The 25-year dispatch for a time-block tender: each 15-minute block of the peak and off-peak
    hours on its own, a shortfall priced at 1.5 x the tariff, and the workbook adding up."""
    import openpyxl

    from engine.sizing import dispatch15 as D

    payload = _payload(_rules(1.5))
    payload["ctx"] = {**payload["ctx"], "trafoLossPct": 0.004, "lineLossPct": 0.005}
    sized = M.solve(payload)
    assert sized["ok"]
    out = D.run(payload, sized["sizes"], sized["tariff"], years=[1, 2], zip_path=tmp_path / "d.zip", xlsx_path=tmp_path / "d.xlsx")
    for y, lp_year in zip(out["years"], [sized["perYear"][0], None]):
        rules = {r["id"]: r for r in y["rules"]}
        assert set(rules) == {"peak", "offpeak"} and all(r["basis"] == "block" and r["penaltyMultiple"] == 1.5 for r in rules.values())
        # within an hour the 15-minute blocks differ (solar ramps), so a block can miss a floor the
        # hour's average met: the 15-minute check is a little stricter than the hourly sizing
        if lp_year:
            hourly = {r["id"]: r for r in lp_year["rules"]}
            for rid in rules:
                assert rules[rid]["hoursMet"] >= hourly[rid]["hoursMet"] - 0.03, (rid, rules[rid], hourly[rid])
            assert y["shortfallMu"] == pytest.approx(sum(r["shortfallMu"] for r in hourly.values()), rel=0.1)
            assert y["deliveredMu"] == pytest.approx(lp_year["deliveredMu"], rel=0.01)
        assert y["shortfallMu"] == pytest.approx(sum(r["shortfallMu"] for r in rules.values()))
    wb = openpyxl.load_workbook(tmp_path / "d.xlsx", read_only=True)
    assert wb.sheetnames == ["About", "Summary", "Monthly", "Year 01", "Year 02"]
    ws = wb["Year 01"]
    rows = list(ws.iter_rows(values_only=True))
    head, body = rows[0], rows[1:]
    assert len(body) == M.HOURS and body[0][:2] == ("01-Jan", "00:00") and body[-1][:2] == ("31-Dec", "23:00")
    col = {name: i for i, name in enumerate(head)}
    y1 = out["years"][0]
    delivered = sum(r[col["Delivered to PPA (MWh)"]] for r in body) / 1000
    short = sum(r[col["Short of the floor (MWh)"]] for r in body) / 1000
    assert delivered == pytest.approx(y1["deliveredMu"], rel=1e-5)
    assert short == pytest.approx(y1["shortfallMu"], rel=1e-4, abs=1e-6)
    assert "Power transformer and line losses (MWh)" in col and "Time-block floor (MWh)" in col
    summary = list(wb["Summary"].iter_rows(values_only=True))
    assert [r[0] for r in summary[1:]] == [1, 2]
    monthly = list(wb["Monthly"].iter_rows(values_only=True))
    assert len(monthly) == 1 + 24
